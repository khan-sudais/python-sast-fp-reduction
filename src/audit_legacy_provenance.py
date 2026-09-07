import argparse
import json
import re
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parent.parent
DEFAULT_OUTPUT = PROJECT_ROOT / "data" / "legacy_provenance_audit.json"

LEGACY_FILES = [
    PROJECT_ROOT / "data" / "aggregated_alert_inventory.xlsx",
    PROJECT_ROOT / "data" / "rq1_cwe_distribution_table.xlsx",
    PROJECT_ROOT / "data" / "rq1_domain_distribution_table.xlsx",
    PROJECT_ROOT / "results_dataset" / "alert_distribution_by_cwe.xlsx",
    PROJECT_ROOT / "results_dataset" / "stratified_ground_truth_sample.xlsx",
    PROJECT_ROOT / "results_main_study" / "main_metrics_by_model_and_variant.xlsx",
    PROJECT_ROOT / "results_main_study" / "statistical_significance_mcnemar.xlsx",
    PROJECT_ROOT / "results_pilot" / "pilot_metrics_table.xlsx",
]

TEXT_ROOTS = [
    PROJECT_ROOT / "README.md",
    PROJECT_ROOT / "src",
    PROJECT_ROOT / "config",
]

TARGET_NUMBERS = {"362", "909", "1029", "300"}


def normalize_name(value):
    return re.sub(r"[^a-z0-9]+", "_", str(value).strip().casefold()).strip("_")


def dataframe_summary(path):
    workbook = pd.ExcelFile(path)
    sheets = {}

    for sheet in workbook.sheet_names:
        dataframe = pd.read_excel(path, sheet_name=sheet)
        normalized_columns = {
            normalize_name(column): str(column)
            for column in dataframe.columns
        }

        unique_counts = {}
        selected_tokens = (
            "repo",
            "repository",
            "scanner",
            "tool",
            "rule",
            "cwe",
            "label",
            "status",
            "alert",
            "finding",
            "file",
            "domain",
            "cohort",
        )

        for normalized, original in normalized_columns.items():
            if any(token in normalized for token in selected_tokens):
                series = dataframe[original]
                unique_counts[original] = {
                    "non_null": int(series.notna().sum()),
                    "unique": int(series.nunique(dropna=True)),
                    "top_values": {
                        str(key): int(value)
                        for key, value in series.astype("string")
                        .fillna("<NA>")
                        .value_counts(dropna=False)
                        .head(15)
                        .items()
                    },
                }

        sheets[sheet] = {
            "rows": int(len(dataframe)),
            "columns": [str(value) for value in dataframe.columns],
            "duplicate_full_rows": int(dataframe.duplicated().sum()),
            "selected_column_profiles": unique_counts,
        }

    return sheets


def find_numbers_in_excel(path):
    workbook = load_workbook(path, read_only=True, data_only=False)
    matches = []

    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                value = cell.value
                if value is None:
                    continue

                text = str(value)
                found = sorted(
                    number
                    for number in TARGET_NUMBERS
                    if re.search(
                        rf"(?<!\d){re.escape(number)}(?!\d)",
                        text,
                    )
                )

                if found:
                    matches.append(
                        {
                            "sheet": worksheet.title,
                            "cell": cell.coordinate,
                            "value": text[:500],
                            "numbers": found,
                        }
                    )

    return matches


def find_numbers_in_text():
    matches = []

    for root in TEXT_ROOTS:
        if not root.exists():
            continue

        paths = [root] if root.is_file() else list(root.rglob("*"))

        for path in paths:
            if not path.is_file():
                continue

            if path.suffix.casefold() not in {
                ".py",
                ".md",
                ".txt",
                ".json",
                ".yaml",
                ".yml",
                ".csv",
            }:
                continue

            try:
                text = path.read_text(encoding="utf-8")
            except UnicodeError:
                try:
                    text = path.read_text(encoding="utf-8-sig")
                except UnicodeError:
                    continue

            for line_number, line in enumerate(text.splitlines(), start=1):
                found = sorted(
                    number
                    for number in TARGET_NUMBERS
                    if re.search(
                        rf"(?<!\d){re.escape(number)}(?!\d)",
                        line,
                    )
                )

                if found:
                    matches.append(
                        {
                            "path": str(path.relative_to(PROJECT_ROOT)),
                            "line": line_number,
                            "text": line[:500],
                            "numbers": found,
                        }
                    )

    return matches


def compare_inventory_and_sample(inventory_path, sample_path):
    inventory = pd.read_excel(inventory_path)
    sample = pd.read_excel(sample_path)

    inventory_lookup = {
        normalize_name(column): str(column)
        for column in inventory.columns
    }
    sample_lookup = {
        normalize_name(column): str(column)
        for column in sample.columns
    }

    common = sorted(set(inventory_lookup).intersection(sample_lookup))
    preferred_tokens = (
        "alert_id",
        "finding_id",
        "id",
        "repository",
        "repo",
        "scanner",
        "tool",
        "rule",
        "cwe",
        "filename",
        "file",
        "line",
        "line_number",
    )

    selected = [
        name
        for name in common
        if any(token == name or token in name for token in preferred_tokens)
    ]

    result = {
        "inventory_rows": int(len(inventory)),
        "sample_rows": int(len(sample)),
        "common_normalized_columns": common,
        "candidate_key_columns": selected,
    }

    if not selected:
        result["sample_rows_matching_inventory_on_candidate_keys"] = None
        result["sample_unique_key_count"] = None
        result["inventory_unique_key_count"] = None
        return result

    inventory_columns = [inventory_lookup[name] for name in selected]
    sample_columns = [sample_lookup[name] for name in selected]

    inventory_keys = set(
        tuple(str(value) for value in row)
        for row in inventory[inventory_columns].fillna("<NA>").itertuples(
            index=False,
            name=None,
        )
    )

    sample_keys = [
        tuple(str(value) for value in row)
        for row in sample[sample_columns].fillna("<NA>").itertuples(
            index=False,
            name=None,
        )
    ]

    result["sample_rows_matching_inventory_on_candidate_keys"] = int(
        sum(key in inventory_keys for key in sample_keys)
    )
    result["sample_unique_key_count"] = int(len(set(sample_keys)))
    result["inventory_unique_key_count"] = int(len(inventory_keys))

    return result


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", default=str(DEFAULT_OUTPUT))
    args = parser.parse_args()

    files = {}
    missing = []

    for path in LEGACY_FILES:
        relative = str(path.relative_to(PROJECT_ROOT))

        if not path.exists():
            missing.append(relative)
            continue

        files[relative] = {
            "sheets": dataframe_summary(path),
            "number_occurrences": find_numbers_in_excel(path),
        }

    inventory_path = PROJECT_ROOT / "data" / "aggregated_alert_inventory.xlsx"
    sample_path = PROJECT_ROOT / "results_dataset" / "stratified_ground_truth_sample.xlsx"

    comparison = None

    if inventory_path.exists() and sample_path.exists():
        comparison = compare_inventory_and_sample(
            inventory_path,
            sample_path,
        )

    result = {
        "target_numbers": sorted(TARGET_NUMBERS, key=int),
        "legacy_files": files,
        "missing_files": missing,
        "text_occurrences": find_numbers_in_text(),
        "inventory_ground_truth_comparison": comparison,
    }

    output = Path(args.output).resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(
        json.dumps(result, indent=2, ensure_ascii=False),
        encoding="utf-8",
    )

    print("=== LEGACY PROVENANCE AUDIT ===")

    for relative, details in files.items():
        print(f"\nFILE: {relative}")

        for sheet, summary in details["sheets"].items():
            print(
                f"  SHEET: {sheet} ROWS: {summary['rows']} "
                f"COLUMNS: {len(summary['columns'])} "
                f"FULL_DUPLICATES: {summary['duplicate_full_rows']}"
            )
            print("  COLUMNS:", " | ".join(summary["columns"]))

            for column, profile in summary[
                "selected_column_profiles"
            ].items():
                top = ", ".join(
                    f"{key}={value}"
                    for key, value in list(
                        profile["top_values"].items()
                    )[:8]
                )
                print(
                    f"  PROFILE {column}: "
                    f"NON_NULL={profile['non_null']} "
                    f"UNIQUE={profile['unique']} "
                    f"TOP=[{top}]"
                )

        occurrences = details["number_occurrences"]

        if occurrences:
            print("  TARGET NUMBER OCCURRENCES:")
            for item in occurrences[:50]:
                print(
                    f"    {item['sheet']}!{item['cell']} "
                    f"{item['numbers']} -> {item['value']}"
                )

    print("\n=== TEXT OCCURRENCES ===")

    for item in result["text_occurrences"]:
        print(
            f"{item['path']}:{item['line']} "
            f"{item['numbers']} -> {item['text']}"
        )

    print("\n=== INVENTORY VS 300-SAMPLE ===")

    if comparison is None:
        print("Comparison unavailable")
    else:
        print(json.dumps(comparison, indent=2))

    print(f"\nAUDIT JSON: {output}")


if __name__ == "__main__":
    main()
